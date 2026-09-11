"""
RAG 服务：解析、分块、嵌入、存储、混合检索（BM25 + 向量）、查询改写
"""
import io
import re
import math
import json
import logging
from collections import defaultdict
from openai import OpenAI
from flask import current_app
from langsmith import traceable
from app.extensions import db
from app.models import Document, DocumentChunk
from app.services.cache_service import cache_service

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {"txt", "md", "pdf", "docx"}


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def parse_file(file_storage):
    filename = file_storage.filename
    ext = filename.rsplit(".", 1)[1].lower()

    if ext == "txt" or ext == "md":
        return file_storage.read().decode("utf-8")

    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_storage.read()))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)

    if ext == "docx":
        from docx import Document as DocxDocument
        doc = DocxDocument(io.BytesIO(file_storage.read()))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    raise ValueError(f"不支持的文件类型: {ext}")


def chunk_text(text, chunk_size=None, overlap=None):
    if chunk_size is None:
        chunk_size = current_app.config["CHUNK_SIZE"]
    if overlap is None:
        overlap = current_app.config["CHUNK_OVERLAP"]

    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:])
            break
        chunk = text[start:end]
        last_sep = -1
        for sep in ["\n\n", "\n", "。", ".", "！", "!", "？", "?", "；", ";", " ", ""]:
            idx = chunk.rfind(sep)
            if idx > chunk_size // 3:
                last_sep = idx + len(sep)
                break
        if last_sep > 0:
            chunk = chunk[:last_sep]
            end = start + last_sep
        chunks.append(chunk.strip())
        start = end - overlap if end > overlap else end

    return [c for c in chunks if c]


def get_embedding_client():
    return OpenAI(
        api_key=current_app.config.get("EMBEDDING_API_KEY", "ollama"),
        base_url=current_app.config["EMBEDDING_BASE_URL"],
    )


def get_embeddings(texts):
    if not texts:
        return []
    client = get_embedding_client()
    model = current_app.config["EMBEDDING_MODEL"]
    response = client.embeddings.create(model=model, input=texts)
    return [item.embedding for item in response.data]


# ============================================================
# BM25 混合检索
# ============================================================

def _tokenize(text):
    """简单的中英文分词"""
    if not text:
        return []
    # 中文：按字切分；英文/数字：按词切分
    tokens = re.findall(r'[\u4e00-\u9fff]|[a-zA-Z0-9]+', text.lower())
    return tokens


def _build_bm25_index(chunks_with_content):
    """构建 BM25 索引"""
    if not chunks_with_content:
        return None, {}

    k1, b = 1.5, 0.75
    n = len(chunks_with_content)
    avg_dl = sum(len(_tokenize(c["content"])) for c in chunks_with_content) / max(n, 1)

    # 文档频率
    df = defaultdict(int)
    doc_tokens = []
    for chunk in chunks_with_content:
        tokens = list(set(_tokenize(chunk["content"])))
        doc_tokens.append(_tokenize(chunk["content"]))
        for t in tokens:
            df[t] += 1

    return {
        "k1": k1, "b": b, "n": n, "avg_dl": avg_dl,
        "df": dict(df), "doc_tokens": doc_tokens,
        "doc_lens": [len(dt) for dt in doc_tokens],
    }, doc_tokens


def _bm25_score(query_tokens, doc_idx, index):
    """计算单个文档的 BM25 分数"""
    score = 0.0
    k1, b, n, avg_dl = index["k1"], index["b"], index["n"], index["avg_dl"]
    doc_tokens = index["doc_tokens"][doc_idx]
    doc_len = index["doc_lens"][doc_idx]

    token_freq = defaultdict(int)
    for t in doc_tokens:
        token_freq[t] += 1

    for qt in query_tokens:
        if qt not in index["df"]:
            continue
        f = token_freq.get(qt, 0)
        if f == 0:
            continue
        idf = math.log((n - index["df"][qt] + 0.5) / (index["df"][qt] + 0.5) + 1)
        tf = (f * (k1 + 1)) / (f + k1 * (1 - b + b * doc_len / max(avg_dl, 1)))
        score += idf * tf

    return score


def _hybrid_search(query, top_k=None):
    """BM25 + 向量混合检索（优化版）"""
    if top_k is None:
        top_k = current_app.config["RAG_TOP_K"]

    vector_weight = current_app.config.get("RAG_VECTOR_WEIGHT", 0.7)
    bm25_weight = current_app.config.get("RAG_BM25_WEIGHT", 0.3)

    # 1. 向量检索（多取一些用于融合）
    query_embedding = get_embeddings([query])[0]
    vector_results = (
        db.session.query(
            DocumentChunk, 
            Document.filename,
            DocumentChunk.embedding.cosine_distance(query_embedding).label('distance')
        )
        .join(Document, DocumentChunk.document_id == Document.id)
        .order_by(DocumentChunk.embedding.cosine_distance(query_embedding))
        .limit(top_k * 3)
        .all()
    )

    if not vector_results:
        return []

    # 2. 如果 BM25 未启用，直接返回向量结果
    if not current_app.config.get("RAG_HYBRID_ENABLED", True):
        return [
            {
                "content": chunk.content,
                "document_id": chunk.document_id,
                "filename": filename,
                "chunk_index": chunk.chunk_index,
                "score": max(0, 1.0 - float(distance)),  # cosine distance → similarity
            }
            for chunk, filename, distance in vector_results[:top_k]
        ]

    # 3. 提取真实的向量相似度分数
    vector_scores = [max(0, 1.0 - float(distance)) for _, _, distance in vector_results]
    
    # 4. BM25 检索
    chunks_data = [
        {
            "content": chunk.content,
            "document_id": chunk.document_id,
            "filename": filename,
            "chunk_index": chunk.chunk_index,
            "idx": i,
        }
        for i, (chunk, filename, _) in enumerate(vector_results)
    ]

    index, _ = _build_bm25_index(chunks_data)
    query_tokens = _tokenize(query)

    # 5. 计算 BM25 分数
    bm25_scores = [_bm25_score(query_tokens, i, index) for i in range(len(vector_results))]
    
    # 6. BM25 分数归一化（最大最小归一化）
    max_bm25 = max(bm25_scores) if bm25_scores else 1.0
    min_bm25 = min(bm25_scores) if bm25_scores else 0.0
    bm25_range = max_bm25 - min_bm25 if max_bm25 > min_bm25 else 1.0
    bm25_norms = [(score - min_bm25) / bm25_range for score in bm25_scores]

    # 7. 计算融合分数
    scored_results = []
    for i, (chunk, filename, _) in enumerate(vector_results):
        hybrid_score = vector_weight * vector_scores[i] + bm25_weight * bm25_norms[i]

        scored_results.append({
            "content": chunk.content,
            "document_id": chunk.document_id,
            "filename": filename,
            "chunk_index": chunk.chunk_index,
            "vector_score": vector_scores[i],
            "bm25_score": bm25_scores[i],
            "score": hybrid_score,
        })

    # 8. 按融合分数排序
    scored_results.sort(key=lambda x: x["score"], reverse=True)
    return scored_results[:top_k]


# ============================================================
# 查询改写
# ============================================================

def _rewrite_query(query):
    """使用 LLM 改写查询，生成多个检索变体以提升召回率"""
    if not current_app.config.get("RAG_QUERY_REWRITE_ENABLED", False):
        return [query]

    try:
        from app.services.ai_service import AIService
        ai = AIService()
        
        # 生成 3 个查询变体（包括原始查询）
        prompt = f"""请将以下用户问题改写为 3 个不同的搜索查询，用于在知识库中检索相关内容。
要求：
1. 保留核心意图
2. 每个查询使用不同的关键词和表达方式
3. 可以扩展相关概念
4. 每行一个查询

原始问题：{query}

搜索查询："""
        
        rewritten = ai.sync_complete(prompt)
        if not rewritten:
            return [query]
        
        # 解析多行结果
        queries = [line.strip() for line in rewritten.strip().split('\n') if line.strip()]
        
        # 确保原始查询在列表中
        if query not in queries:
            queries.insert(0, query)
        
        # 限制最多 3 个查询
        return queries[:3]
        
    except Exception as e:
        logger.warning("查询改写失败: %s", e)
        return [query]


def _rerank_results(query, results, top_k=None):
    """使用 LLM 对检索结果进行重排序"""
    if top_k is None:
        top_k = current_app.config["RAG_TOP_K"]
    
    if not current_app.config.get("RAG_RERANK_ENABLED", False):
        return results[:top_k]
    
    if len(results) <= top_k:
        return results
    
    try:
        from app.services.ai_service import AIService
        ai = AIService()
        
        # 构建候选列表
        candidates = []
        for i, r in enumerate(results[:top_k * 3], 1):
            preview = r['content'][:100] + '...' if len(r['content']) > 100 else r['content']
            candidates.append(f"[{i}] {preview}")
        
        prompt = f"""根据以下用户问题，对检索结果按相关性重新排序。
只返回排序后的编号列表（如：3,1,4,2），不要解释。

用户问题：{query}

检索结果：
{chr(10).join(candidates)}

排序后的编号（从最相关到最不相关）："""
        
        reranked = ai.sync_complete(prompt)
        if not reranked:
            return results[:top_k]
        
        # 解析排序结果
        order = [int(x.strip()) - 1 for x in reranked.replace('，', ',').split(',') if x.strip().isdigit()]
        
        # 按新顺序重排
        reranked_results = []
        for idx in order:
            if 0 <= idx < len(results):
                reranked_results.append(results[idx])
        
        # 补充未被排序的结果
        used_indices = set(order)
        for i, r in enumerate(results):
            if i not in used_indices and len(reranked_results) < top_k:
                reranked_results.append(r)
        
        return reranked_results[:top_k]
        
    except Exception as e:
        logger.warning("重排序失败: %s", e)
        return results[:top_k]


# ============================================================
# 公共接口
# ============================================================

@traceable(name="rag_process_and_store")
def process_and_store(file_storage):
    filename = file_storage.filename
    ext = filename.rsplit(".", 1)[1].lower()
    file_size = file_storage.content_length or 0

    content = parse_file(file_storage)
    chunks = chunk_text(content)

    if not chunks:
        raise ValueError("文件中没有提取到有效内容")

    embeddings = get_embeddings(chunks)

    doc = Document(
        filename=filename,
        file_type=ext,
        file_size=file_size,
        chunk_count=len(chunks),
    )
    db.session.add(doc)
    db.session.flush()

    for i, (chunk, embedding) in enumerate(zip(chunks, embeddings)):
        dc = DocumentChunk(
            document_id=doc.id,
            chunk_index=i,
            content=chunk,
            embedding=embedding,
        )
        db.session.add(dc)

    db.session.commit()
    return doc


@traceable(name="rag_search")
def search(query, top_k=None):
    """搜索知识库（支持缓存 + 多查询 + 混合检索 + 重排序）"""
    if top_k is None:
        top_k = current_app.config["RAG_TOP_K"]

    # 0. 检查缓存（相同查询直接返回）
    if current_app.config.get("CACHE_ENABLED", True):
        cached = cache_service.get(query, f"rag_{top_k}")
        if cached is not None:
            logger.info(f"RAG cache hit: {query[:50]}")
            return cached

    # 1. 查询改写（生成多个查询变体）
    queries = _rewrite_query(query)

    # 2. 对每个查询变体执行混合检索并合并去重
    all_results = {}
    for q in queries:
        results = _hybrid_search(q, top_k=top_k * 2)  # 多取一些用于合并
        for r in results:
            key = (r['document_id'], r['chunk_index'])
            if key not in all_results or all_results[key]['score'] < r['score']:
                all_results[key] = r

    # 3. 按分数排序
    merged_results = sorted(all_results.values(), key=lambda x: x['score'], reverse=True)

    # 4. 重排序（可选）
    final_results = _rerank_results(query, merged_results, top_k=top_k)

    # 5. 存入缓存
    if current_app.config.get("CACHE_ENABLED", True):
        cache_service.set(query, final_results, f"rag_{top_k}", ttl_hours=1)

    return final_results


def delete_document(doc_id):
    doc = db.session.get(Document, doc_id)
    if not doc:
        return False
    db.session.delete(doc)
    db.session.commit()
    return True


def list_documents():
    docs = Document.query.order_by(Document.created_at.desc()).all()
    return [
        {
            "id": d.id,
            "filename": d.filename,
            "file_type": d.file_type,
            "file_size": d.file_size,
            "chunk_count": d.chunk_count,
            "created_at": d.created_at.isoformat(),
        }
        for d in docs
    ]
